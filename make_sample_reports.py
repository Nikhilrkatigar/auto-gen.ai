"""
make_sample_reports.py - generate synthetic concurrent audit reports (.docx).

The real reports contain customer data and stay private. These files copy
their layout (annexure headings + tables, daily / weekly / monthly formats)
with made-up branches, people, accounts and amounts, so the RAG pipeline can
be run and evaluated by anyone.

Usage::

    python make_sample_reports.py          # writes sample_reports/
"""
import random
from pathlib import Path

import docx

OUT = Path("sample_reports")
BRANCHES = {"1101": "LAKEVIEW", "1102": "HILLTOP", "1103": "GREENFIELD", "1104": "MARKET YARD"}
FIRST = ["RAVI", "ANITA", "SURESH", "KAVYA", "MOHAN", "PRIYA", "ARJUN", "LATHA", "VINOD", "DEEPA", "KIRAN", "MEERA"]
LAST = ["RAO", "SHETTY", "KULKARNI", "PATIL", "NAIK", "HEGDE", "JOSHI", "DESAI", "BHAT", "REDDY"]
LOANS = [("CC General", 11.40), ("Micro Enterprises - Mfg", 11.75), ("DL Agst NSC Personal", 13.00),
         ("CC SHG NRLM Agri", 7.00), ("Business Loan OD", 9.60), ("Education Loan Guar", 10.40),
         ("Housing Loan Salaried", 8.50), ("Gold Loan Agri", 9.00)]
BGL_HEADS = ["NL CASH GIVEN TO SERVICE PROVIDER", "PPF SCHEME REMITTED", "SENIOR CITIZEN DEP SCH REMITTED",
             "NL DEPOSIT WITH ELECTRICITY BOARD", "NL ADVANCE RENT-LANDLORD", "INT SUBVENTION AGRICULTURE",
             "SUNDRY CREDITORS - CLEARING", "TDS PAYABLE ACCOUNT"]
CARD_TYPES = ["SB GEN PUB IND URBAN", "NCMC Contactless Card", "CD SOLE PROP", "SB SMALL ACCOUNT"]


def person(rng):
    title = rng.choice(["Mr.", "Mrs.", "Miss"])
    return f"{title} {rng.choice(FIRST)} {rng.choice(LAST)}"


def acct(rng, n=10):
    return str(rng.randint(10 ** (n - 1), 10 ** n - 1))


def money(value):
    whole = f"{int(value):d}"
    head, tail = whole[:-3], whole[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return ",".join(groups + [tail]) + f".{int(round(value * 100)) % 100:02d}"


def table(d, header, rows):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)


def date(rng, start=1, end=28):
    return f"{rng.randint(start, end):02d}/05/2026"


def interest_rows(rng, n):
    rows = []
    for _ in range(n):
        name, rate = rng.choice(LOANS)
        delta = rng.choice([-1, 1]) * rng.choice([0.15, 0.5, 1.0, 3.0])
        rows.append([acct(rng), name, f"{rate + delta:.2f}", f"{rate:.2f}"])
    return rows


def staff_rows(rng, n, dated=True):
    rows = []
    for _ in range(n):
        amount = rng.choice([50000, 74000, 100000, 225000, 360000, 400000])
        row = [acct(rng), f"{rng.choice(FIRST)} {rng.choice(LAST)}", ".00", f"-{amount}.00"]
        rows.append([date(rng)] + row if dated else row)
    return rows


def daily(rng, code, branch, i):
    d = docx.Document()
    d.add_paragraph("DAILY CONCURRENT AUDIT REPORT")
    d.add_paragraph("Periodicity: Daily      Concurrent Audit Format No – D-1")
    d.add_paragraph(f"Branch: {branch}    Region: North Region   Date of report: 15/05/2026")
    table(d, ["S N", "Report on", "Irregularity Observed (Yes/ No)"],
          [[1, "Transaction of work available to the field in B/S & P/L.", "YES"],
           [2, "Transactions from nominal to loan/ other accounts.", "NO"],
           [6, "Manual tinkering of interest spread in any loan heads.", "YES"],
           [15, "Individual transactions of ₹45000/- and above in all staff accounts.", "YES"]])
    d.add_paragraph("Annexure-1 Transaction of work available to the field in B/S & P/L (Nominal/ BGL accounts)")
    table(d, ["A/c No", "A/c Name", "Credit Amount", "Debit Amount"],
          [[acct(rng, 13), h, "", money(rng.randint(5_000, 5_00_00_000))] for h in rng.sample(BGL_HEADS, 4)])
    d.add_paragraph("Annexure-6 Manual tinkering of interest spread in any loan heads. (CCOD)")
    table(d, ["Account No", "Account Description", "Rate as per Account", "Rate as per Product"], interest_rows(rng, 4))
    d.add_paragraph("(DL/TL)")
    table(d, ["Account No", "Account Description", "Rate as per Account", "Rate as per Product"], interest_rows(rng, 3))
    d.add_paragraph("Annexure-11 Reconciliation of entries of ANINA & Other Liabilities (Greater than 30 Days reported)")
    table(d, ["A/c No", "Name", "Amt", "Age in Days"],
          [[acct(rng, 11), h, money(rng.randint(1_000, 3_00_000)), rng.randint(31, 3600)] for h in rng.sample(BGL_HEADS, 3)])
    d.add_paragraph("Annexure-15 Individual transactions of an amount of ₹45000/- and above in all staff accounts.")
    table(d, ["A/c No", "Name", "Credit", "Debit"], staff_rows(rng, 2, dated=False))
    d.add_paragraph("Annexure-17 Opening of new SB/ CD accounts.")
    table(d, ["Date", "A/c No", "Name", "Type"], [[date(rng), acct(rng), person(rng), rng.choice(CARD_TYPES)] for _ in range(4)])
    d.save(OUT / f"Daily AR - {i}.{code} - {branch}.docx")


def weekly(rng, code, branch, i):
    d = docx.Document()
    d.add_paragraph("CONSOLIDATED REPORT OF IRREGULARITY OBSERVED IN OTHER PARAMETERS")
    d.add_paragraph(f"Branch Name: {branch}")
    d.add_paragraph("Weekly Report for the period 02/05/2026 to 09/05/2026      Date of report: 09/05/2026")
    table(d, ["S N", "Report on", "Date of TXN", "Irregularities/ Observation/ Remark"],
          [[1, "Interest variation Report.", "02-05-2026 to 09-05-2026", "NIL"],
           [2, "Manual tinkering of interest spread in loan heads.", "02-05-2026 to 09-05-2026", "Refer Annexure-2"],
           [3, "Linking of CIFs/ Delinking of CIFs.", "02-05-2026 to 09-05-2026", "Refer Annexure-3"]])
    d.add_paragraph("Annexure-2 Manual tinkering of interest spread in any loan heads (CCOD)")
    table(d, ["A/c No", "Name", "Rate as per A/c", "Rate as per Product"], interest_rows(rng, 4))
    d.add_paragraph("Annexure-3 Linking of CIFs/ Delinking of CIFs.")
    table(d, ["Date", "A/c No", "CIF No", "Name", "Type"],
          [[date(rng, 2, 9), acct(rng), acct(rng), person(rng), rng.choice(["CIF LINKING", "CIF DELINKING"])] for _ in range(2)])
    d.add_paragraph("Annexure-9 Change of mobile number in CIF.")
    table(d, ["Date", "CIF No.", "Name", "Old Mobile No.", "New Mobile No."],
          [[date(rng, 2, 9), acct(rng), person(rng), "91" + acct(rng), "91" + acct(rng)] for _ in range(3)])
    d.add_paragraph("Annexure-11 Individual transactions of an amount of ₹45000/- and above in all staff accounts.")
    table(d, ["Date", "A/c No", "Name", "Credit", "Debit"], staff_rows(rng, 2))
    d.add_paragraph("Annexure-12 Report accounts becoming operative from Inoperative")
    table(d, ["Date", "A/c No", "Name", "DATE OF INOPERATIVE", "DATE OF OPERATIVE"],
          [[date(rng, 2, 9), acct(rng), person(rng), f"{rng.randint(1, 28):02d}/0{rng.randint(1, 9)}/2024", date(rng, 2, 9)] for _ in range(2)])
    d.save(OUT / f"Weekly AR - {i}. {code} - {branch}.docx")


def monthly(rng, code, branch, i):
    d = docx.Document()
    obs = "TO BE CORRECTED BY BRANCH"
    for title, header, rows in [
        ("REPORT ON INTEREST VARIATION IN LOAN ACCOUNTS",
         ["S N", "Account No.", "Account Name", "ROI as per CBS", "ROI as per Product", "Audit Observations"],
         [[n] + r + [obs] for n, r in enumerate(interest_rows(rng, 5), 1)]),
        ("REPORT ON TRANSACTIONS IN ACCOUNTS OF STAFF",
         ["S N", "Transaction Date", "Account No.", "Account Name", "Transaction Amount", "Transaction type (Credit/ Debit)", "Audit Observations"],
         [[n, date(rng), acct(rng), f"{rng.choice(FIRST)} {rng.choice(LAST)}", f"{rng.choice([60000, 93274.53, 225000, 360000])}",
           rng.choice(["Credit", "Debit"]), "Branch to ensure source of funds"] for n in range(1, 4)]),
        ("REPORT ON RECONCILIATION OF ENTRIES OF ANINA & OTHER LIABILITIES",
         ["S N", "Transaction Date", "Account No.", "Account Name", "Transaction Amount", "Transaction Type (Debit/ Credit)", "Audit Observations"],
         [[n, f"{rng.randint(1, 28):02d}-0{rng.randint(1, 9)}-2{rng.randint(0, 5)}", acct(rng, 11), h, money(rng.randint(1_000, 10_00_000)),
           "Debit", "TO BE ADJUSTED BY BRANCH"] for n, h in enumerate(rng.sample(BGL_HEADS, 3), 1)]),
        ("REPORT ON OPENING OF NEW SB/ CD ACCOUNTS",
         ["S N", "Transaction Date", "Account No.", "Account Name", "KYC complied (Yes/ No)", "Audit Observations"],
         [[n, date(rng), acct(rng), person(rng), rng.choice(["Yes", "No"]),
           "However, we were unable to verify the scanned copies of the KYC documents in the portal in certain cases."
           if branch == "GREENFIELD" else "KYC verified"] for n in range(1, 4)]),
    ]:
        d.add_paragraph("Periodicity: Monthly  Concurrent Audit Format No. M-NL")
        d.add_paragraph(title)
        d.add_paragraph(f"Branch: {branch}  Region: North Region  Reporting Month: MAY - 2026")
        table(d, header, rows)
    d.save(OUT / f"Monthly Report May - {i}.{code} - {branch}.docx")


def main():
    rng = random.Random(2026)  # fixed seed: same reports every run
    OUT.mkdir(exist_ok=True)
    for i, (code, branch) in enumerate(BRANCHES.items(), 1):
        daily(rng, code, branch, i)
        weekly(rng, code, branch, i)
        monthly(rng, code, branch, i)
    print(f"Wrote {len(list(OUT.glob('*.docx')))} synthetic reports to {OUT}/")


if __name__ == "__main__":
    main()
