"""Tests for rag.py that run offline: no embedding model and no API calls."""
from types import SimpleNamespace

import numpy as np

import rag
from rag import BM25, Chunk, Index


def make_chunk(text, branch="LAKEVIEW"):
    return Chunk(text, "Weekly AR - LAKEVIEW.docx", branch, "Weekly", "", "Annexure-11")


class FakeEmbedder:
    """Bag-of-letters vectors: enough to exercise the dense path."""

    def _vec(self, text):
        v = np.zeros(26, dtype="float32")
        for ch in text.lower():
            if "a" <= ch <= "z":
                v[ord(ch) - 97] += 1
        return v / (np.linalg.norm(v) or 1)

    def docs(self, texts):
        return np.array([self._vec(t) for t in texts])

    def query(self, text):
        return self._vec(text)


def test_mask_numbers_keeps_last_four_digits():
    assert rag.mask_numbers("A/c 5520031874 mobile 919000012345") == "A/c XXXXXX1874 mobile XXXXXXXX2345"
    assert rag.mask_numbers("Amount 4,00,000.00 rate 16.00") == "Amount 4,00,000.00 rate 16.00"


def test_initials_drop_honorifics():
    assert rag.initials("Mr. RAVI KUMAR SHARMA") == "R. K. S."
    assert rag.initials("Mrs. ANITA .") == "A."


def test_expand_amounts():
    assert "360000" in rag.expand_amounts("credit of 3.6 lakh")
    assert "10000000" in rag.expand_amounts("1 crore")


def test_bm25_ranks_matching_document_first():
    scores = BM25(["staff account debit", "interest rate variation", "new SB account"]).scores("interest rate")
    assert scores.index(max(scores)) == 1


def test_search_filters_by_branch_named_in_question():
    chunks = [make_chunk("staff debit 100000", "LAKEVIEW"), make_chunk("staff debit 100000", "HILLTOP")]
    index = Index.build(chunks, FakeEmbedder())
    hits = index.search("staff debit at HILLTOP", k=5)
    assert [c.branch for c, _ in hits] == ["HILLTOP"]


def test_chunk_report_masks_people_and_numbers(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Annexure-11 Individual transactions in all staff accounts.")
    t = d.add_table(rows=2, cols=3)
    for cell, text in zip(t.rows[0].cells, ["A/c No", "Name", "Debit"]):
        cell.text = text
    for cell, text in zip(t.rows[1].cells, ["4410029387", "ANAND RAO", "-100000.00"]):
        cell.text = text
    path = tmp_path / "Weekly AR - 3. 1103 - LAKEVIEW.docx"
    d.save(path)

    [chunk] = rag.chunk_report(path)
    assert chunk.branch == "LAKEVIEW" and chunk.report_type == "Weekly"
    assert "XXXXXX9387" in chunk.text and "A. R." in chunk.text
    assert "4410029387" not in chunk.text and "ANAND" not in chunk.text


def test_answer_sends_sources_and_reads_text():
    sent = {}

    def create(**kwargs):
        sent.update(kwargs)
        return SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text="Debit of 1,00,000 [1]")])

    client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(create=create)))
    out = rag.answer("staff debit?", [(make_chunk("Debit: -100000.00"), 1.0)], client=client)
    assert out == "Debit of 1,00,000 [1]"
    assert sent["model"] == rag.CLAUDE_MODEL
    assert "[1] (Weekly AR - LAKEVIEW.docx)" in sent["messages"][0]["content"]


def test_answer_handles_refusal():
    client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(
        create=lambda **_: SimpleNamespace(stop_reason="refusal", content=[]))))
    assert "declined" in rag.answer("q", [], client=client)
